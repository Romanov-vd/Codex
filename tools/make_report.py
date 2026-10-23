"""Create the common report for variants 1 and 2; keep the sample typography."""
from pathlib import Path
import json
import math
import shutil
import subprocess
import tempfile
import numpy as np
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'reports'
FIG=OUT/'figures'
FIG.mkdir(parents=True,exist_ok=True)
metrics={}
for name in ['shapes','detail']:
    clean=np.asarray(Image.open(ROOT/'test_input'/f'{name}.png'),dtype=float)
    noisy=np.asarray(Image.open(ROOT/'test_input'/f'{name}_noisy.png'),dtype=float)
    methods={}
    for method in ['noisy','mean3','mean5','gauss5']:
        image=noisy if method=='noisy' else np.asarray(Image.open(ROOT/'test_output'/f'{name}_noisy'/f'{method}.png'),dtype=float)
        mse=float(np.mean((image-clean)**2))
        methods[method]={'mse':round(mse,2),'psnr':round(10*math.log10(255**2/mse),2),
                         'bias':round(float((image-clean).mean()),2)}
    metrics[name]={'noise_mean':float((noisy-clean).mean()),
                   'noise_std':float((noisy-clean).std()),'methods':methods}
(OUT/'metrics.json').write_text(json.dumps(metrics,indent=2),encoding='utf-8')
font=ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',15)
def montage(name,items,columns):
    canvas=Image.new('RGB',(272*columns,224*math.ceil(len(items)/columns)),'white')
    draw=ImageDraw.Draw(canvas)
    for i,(label,path) in enumerate(items):
        x,y=(i%columns)*272,(i//columns)*224
        canvas.paste(Image.open(path).convert('RGB'),(x+8,y+28))
        draw.text((x+8,y+5),label,font=font,fill='black')
    path=FIG/(name+'.png')
    canvas.save(path)
    return path
inputs=montage('inputs',[(label,ROOT/'test_input'/name) for label,name in [
    ('А: фигуры, оригинал','shapes.png'),('А: гауссов шум','shapes_noisy.png'),
    ('Б: детали, оригинал','detail.png'),('Б: гауссов шум','detail_noisy.png')]],2)
edges=[]
averages=[]
for scene,label in [('shapes','А'),('detail','Б')]:
    for method,title in [('gray','Исходное'),('laplace','Лаплас'),('roberts','Робертс'),('sobel','Собел')]:
        edges.append((label+': '+title,ROOT/'test_output'/scene/(method+'.png')))
    for method,title in [('gray','С шумом'),('mean3','Среднее 3×3'),('mean5','Среднее 5×5'),('gauss5','Гаусс 5×5')]:
        averages.append((label+': '+title,ROOT/'test_output'/(scene+'_noisy')/(method+'.png')))
edge_figure=montage('edges',edges,4)
noise_edges=montage('edges_noisy',[(title,ROOT/'test_output'/'shapes_noisy'/(method+'.png')) for method,title in [
    ('gray','Вход с шумом'),('laplace','Лаплас'),('roberts','Робертс'),('sobel','Собел')]],4)
average_figure=montage('averages',averages,4)

doc=Document()
for style in doc.styles:
    if hasattr(style,'font'):
        style.font.name='Times New Roman'
        style.font.color.rgb=RGBColor(0,0,0)
        style.font.underline=False
for color in doc.styles.element.iter(qn('w:color')):
    color.attrib.clear()
    color.set(qn('w:val'),'000000')
section=doc.sections[0]
section.different_first_page_header_footer=True
footer=section.footer.paragraphs[0]
footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
field=OxmlElement('w:fldSimple')
field.set(qn('w:instr'),'PAGE')
footer._p.append(field)
def p(text,center=False):
    paragraph=doc.add_paragraph(text)
    if center:paragraph.alignment=WD_ALIGN_PARAGRAPH.CENTER
    return paragraph
def h(text):doc.add_heading(text,level=1)
def subheading(text):
    paragraph=p(text,True)
    for run in paragraph.runs:run.bold=True
def algorithm(title,steps):
    p(title)
    for i,step in enumerate(steps,1):p(str(i)+'. '+step)
def eq(text):p(text,True)
def figure(path,caption,width=16.5):
    doc.add_picture(str(path),width=Cm(width))
    p(caption,True)
def table(headers,rows):
    t=doc.add_table(rows=1,cols=len(headers))
    t.style='Table Grid'
    for cell,value in zip(t.rows[0].cells,headers):cell.text=value
    for values in rows:
        for cell,value in zip(t.add_row().cells,values):cell.text=str(value)

p('МИНОБРНАУКИ РОССИИ',True)
p('Воронежский государственный университет\nФакультет прикладной математики, информатики и механики\nКафедра математического обеспечения ЭВМ',True)
p('\n\n\n')
p('Цифровая обработка изображений',True)
p('Лабораторная работа № 2',True)
p('Выделение контуров и усредняющая фильтрация\nВарианты 1 и 2',True)
p('\n\n\n')
p('Выполнил(а): __________________________\nГруппа: ______________________________\nПреподаватель: Сырых А. С.')
p('\n\n')
p('Воронеж, 2026',True)
doc.add_page_break()

h('Постановка задачи')
p('Необходимо реализовать на языке C++ обработку полутоновых изображений по двум выбранным вариантам задания.')
p('Вариант 1 — выделение контуров с помощью операторов Лапласа, Робертса и Собела.')
p('Вариант 2 — подавление гауссова шума пространственными усредняющими фильтрами. Зашумлённые изображения следует подготовить в графическом редакторе.')
p('Перед выполнением обработки требуется завершить заданный класс Image: обеспечить чтение и сохранение PNG, доступ к пикселям, поддержку одного–четырёх каналов и независимое копирование изображений. Расширять интерфейс класса новыми методами не допускается. Работа должна сопровождаться тестовыми изображениями, которые показывают различия рассматриваемых методов.')

h('Математическое описание алгоритма')
p('Пусть f(x,y) задаёт яркость пикселя в позиции (x,y), а W и H — размеры изображения. Для восьмибитного представления допустимы значения от 0 до 255. Если исходное изображение цветное, перед обработкой вычисляется полутоновая яркость по формуле (1). Для одно- и двухканального входа используется существующий канал яркости.')
eq('f(x,y) = round(0.299R + 0.587G + 0.114B).  (1)')
p('При обращении за пределы растра берётся ближайший допустимый пиксель. Такое продолжение сохраняет постоянный фон у границ. Каждый метод читает исходный массив и формирует отдельное выходное изображение.')

h('Алгоритмы выделения контуров')
p('Граница объекта проявляется как быстрое изменение яркости. Для её обнаружения используются дискретные приближения производных: первые разности у Робертса и Собела и вторая разность у Лапласа.')
subheading('Оператор Собела')
p('В окне 3×3 рассчитываются изменения яркости по горизонтали и вертикали. Маски (2) объединяют разностную оценку с локальным сглаживанием в поперечном направлении:')
eq('Sx = [−1 0 1; −2 0 2; −1 0 1];\nSy = [−1 −2 −1; 0 0 0; 1 2 1].  (2)')
p('Компоненты градиента и его амплитуда определяются по (3). Суммирование ведётся по смещениям i,j = −1,0,1 относительно текущего пикселя:')
eq('Gx = Σ_i Σ_j Sx[j,i] f(x+i,y+j);\nGy = Σ_i Σ_j Sy[j,i] f(x+i,y+j);\nM(x,y) = √(Gx²+Gy²).  (3)')
p('Для записи в восьмибитный канал результат ограничивается и округляется по (4). Эта операция не является растяжением яркостей по максимуму конкретного изображения. Она также не выполняет пороговую бинаризацию:')
eq('Q(t) = round(min(255, max(0,t))); g(x,y) = Q(M(x,y)).  (4)')
algorithm('Алгоритм 1. Оператор Собела',[
    'Получить полутоновый вход по формуле (1).',
    'Для каждого пикселя найти Gx и Gy с использованием масок (2).',
    'Вычислить амплитуду по (3) и подготовить значение канала по (4).',
    'Поместить значение в отдельное результирующее изображение.'])
subheading('Оператор Робертса')
p('Вместо окна 3×3 используются две диагонали ячейки 2×2. Соответствующие маски записываются в виде (5):')
eq('Rx = [1 0; 0 −1]; Ry = [0 1; −1 0].  (5)')
p('Из этих масок следуют две разности (6). Амплитуда получается как корень из суммы их квадратов, после чего применяется ограничение (4):')
eq('a = f(x,y)−f(x+1,y+1);\nb = f(x+1,y)−f(x,y+1); M(x,y) = √(a²+b²).  (6)')
algorithm('Алгоритм 2. Оператор Робертса',[
    'Сформировать яркостное изображение по (1).',
    'В каждом окне 2×2 вычислить диагональные разности по (6).',
    'Найти амплитуду отклика и записать её с учётом (4).'])
subheading('Оператор Лапласа')
p('В непрерывном представлении Лапласиан равен сумме вторых производных по двум пространственным координатам:')
eq('Δf = ∂²f/∂x² + ∂²f/∂y².  (7)')
p('В работе применяется четырёхсвязное дискретное приближение с ядром (8):')
eq('K = [0 1 0; 1 −4 1; 0 1 0].  (8)')
p('Поскольку вторая разность может иметь любой знак, для отображения используется её модуль. Формула расчёта приведена в (9):')
eq('L = f(x−1,y)+f(x+1,y)+f(x,y−1)+f(x,y+1)−4f(x,y);\ng(x,y) = Q(|L|).  (9)')
algorithm('Алгоритм 3. Оператор Лапласа',[
    'Преобразовать вход в полутоновое представление по (1).',
    'Вычислить вторую разность с использованием маски (8).',
    'Взять модуль отклика и сформировать результат по (9).'])

h('Линейные усредняющие фильтры')
p('Для уменьшения случайных колебаний яркости выходной пиксель вычисляется по его окрестности. При этом сглаживается не только шум, но и полезные изменения сигнала, поэтому выбор окна влияет на сохранность границ и мелких объектов.')
subheading('Прямоугольный усредняющий фильтр (Box Blur)')
p('Все значения квадратного окна имеют одинаковый вес. Для нечётного размера k радиус r=(k−1)/2, а число элементов равно k². Арифметическое усреднение описывается формулой (10):')
eq('g(x,y) = Q((1/k²) Σ_i Σ_j f(x+i,y+j)), i,j = −r,…,r.  (10)')
algorithm('Алгоритм 4. Прямоугольное усреднение',[
    'Выбрать нечётный размер окна и определить его радиус.',
    'Присвоить каждому из k² элементов вес 1/k².',
    'Вычислить сумму значений окрестности с этими весами по (10).',
    'Округлить результат и сохранить его в новом изображении.'])
subheading('Фильтр Гаусса (Gaussian Blur)')
p('В отличие от равномерного окна, здесь центральные пиксели сильнее влияют на выход. Распределение весов задаётся гауссовой функцией (11), где σ определяет пространственный масштаб сглаживания:')
eq('G(i,j) = (1/(2πσ²)) exp(−(i²+j²)/(2σ²)).  (11)')
p('На конечном окне веса необходимо нормировать по (12). Общий множитель 1/(2πσ²) при делении сокращается, поэтому реализация вычисляет только экспоненциальную часть:')
eq('w(i,j) = G(i,j) / Σ_i Σ_j G(i,j); Σ_i Σ_j w(i,j) = 1.  (12)')
p('После подготовки ядра выход определяется взвешенной суммой (13):')
eq('g(x,y) = Q(Σ_i Σ_j w(i,j) f(x+i,y+j)).  (13)')
algorithm('Алгоритм 5. Гауссово усреднение',[
    'Задать размер окна k и положительное значение σ.',
    'Определить веса по (11) и нормировать их согласно (12).',
    'Для каждого пикселя вычислить взвешенную сумму по (13).',
    'Сформировать отдельный выходной массив с учётом (4).'])

h('Программное описание алгоритма')
p('Программа написана на C++17. Объект Image отвечает за память, PNG и доступ к пикселям. Его исходный интерфейс сохранён; алгоритмы вынесены в пространство имён Filters. Данные хранятся в std::unique_ptr<uint8_t[]>, а копирование создаёт независимый массив. Для серого входа getColor возвращает яркость в r, устанавливая g и b в ноль; setColor сохраняет альфа-канал. При недопустимых координатах доступ безопасно игнорируется.')
p('Функция Filters::toGrayscale формирует одно яркостное изображение по (1). Полученный объект используется при обработке, а исходные данные не изменяются. Вспомогательная sample повторяет ближайший граничный пиксель, put реализует ограничение и округление по (4).')
p('Filters::applySobel принимает изображение по константной ссылке, получает его полутоновое представление и вызывает общий внутренний расчёт edges с выбором Собела. В нём выполняется алгоритм 1: накапливаются компоненты по маскам (2), затем вычисляется амплитуда (3). Соответствующие фрагменты кода приведены в приложении А.')
p('Filters::applyRoberts использует тот же внутренний расчёт с другим режимом. Для каждой позиции берутся две диагональные разности по (6); их амплитуда преобразуется в допустимый диапазон. Это реализация алгоритма 2, представленная в приложении А.')
p('Filters::applyLaplacian выбирает четырёхсвязную вторую разность и её модуль согласно (8), (9). Функция соответствует алгоритму 3; её реализация и общий расчёт edges также приведены в приложении А.')
p('Filters::applyBoxBlur получает изображение и размер окна. Через внутреннюю функцию average она задаёт равные веса и выполняет алгоритм 4 по (10). Filters::applyGaussianBlur дополнительно проверяет положительность и конечность σ, формирует нормированные гауссовы веса и выполняет алгоритм 5 по (11)–(13). Эти фрагменты приведены в приложении Б.')
p('Точка входа находится в Source.cpp. main выбирает вход и каталог результата, а runFilters последовательно запускает пять функций обработки: три контурные и две усредняющие. Арифметическое среднее применяется с окнами 3×3 и 5×5, Гаусс — с окном 5×5 при σ=1. Вместе с полутоновым входом записываются семь выходных PNG. Сохранение проверяет saveImage; стандартные исключения перехватываются в main, сопровождаются сообщением и кодом завершения 1.')
p('Без аргументов используется готовый файл test_input/shapes_noisy.png. Можно передать два пути для своего изображения либо --random с необязательным seed для дополнительной случайной сцены. Она состоит из градиента, прямоугольников и тонких штрихов, имеет размер 320×224 и шум со стандартным отклонением 18. Если готовый пример отсутствует, такой тест создаётся автоматически. Этот режим не используется для численных таблиц отчёта: основные шумные входы подготовлены в ImageMagick согласно условию варианта 2.')
p('В решении Visual Studio предусмотрены отдельные проекты Lab и Tests. Проверка содержит 18 тестовых случаев и 323 условия, все пройдены. Проверяются загрузка и сохранение, каналы, прозрачность, копирование, граничный доступ и заранее рассчитанные ответы фильтров. Например, на ступени 0→40 ответы Лапласа, Робертса и Собела равны соответственно 40, 57 и 160; среднее 3×3 для одиночного импульса 255 даёт 28.')

h('Результат работы программы')
p('Для эксперимента подготовлены два полутоновых изображения размером 256×192. В наборе А присутствуют крупные фигуры, ровный фон и плавное изменение яркости. В наборе Б размещены тонкие штрихи, клетки, окружность и диагональ. Первый набор показывает изменение шума на широких областях, второй позволяет обнаружить потерю небольших деталей. Для каждой сцены сохранены чистый эталон и шумная версия.')
figure(inputs,'Рисунок 1. Слева — чистые изображения А и Б, справа — соответствующие версии с гауссовым шумом.',15)
p('Шум добавлен вне C++-программы редактором ImageMagick с параметрами +noise Gaussian и -attenuate 0.9, с последующим сохранением в восьмибитный PNG. Начальные значения генератора равны 731 и 1201. Возможность применять другой графический редактор прямо предусмотрена вариантом 2.')
p(f"Измеренное стандартное отклонение разности с чистым эталоном равно {metrics['shapes']['noise_std']:.2f} для А и {metrics['detail']['noise_std']:.2f} для Б. Это характеристика шума в уровнях яркости; она не совпадает по смыслу с σ=1 пространственного гауссова фильтра.")
figure(edge_figure,'Рисунок 2. Контуры чистых сцен: сверху А, снизу Б. В каждом ряду — вход, Лаплас, Робертс и Собел.')
p('Все три оператора реагируют на границы крупных фигур. У Лапласа после взятия модуля становятся видимыми отклики разных знаков с обеих сторон перепада, что объясняет двойную границу. Робертс даёт компактный отклик на диагональных разностях, а Собел формирует более широкий и яркий контур из-за маски 3×3 и её коэффициентов. Поэтому сравнение только по яркости линии недостаточно.')
p('На сцене с деталями рядом расположенные границы штрихов и клеток могут сливаться. Внутри линейного градиента вторая разность близка к нулю, тогда как первые разности реагируют на плавное изменение яркости. Остаточные слабые отклики связаны в том числе с округлением значений при создании градиента.')
figure(noise_edges,'Рисунок 3. Шумный вход А и результаты контурных операторов в том же порядке.')
p('При обработке шумного входа появляются ложные контуры на исходно однородных участках. Случайные отклонения также являются перепадами яркости и усиливаются разностными операциями. Поперечное сглаживание Собела снижает влияние отдельных колебаний, но не устраняет его полностью. Значения выше 255 насыщаются по (4), поэтому после ограничения яркость уже не отражает точную величину производной.')
figure(average_figure,'Рисунок 4. Подавление шума: вход, среднее 3×3, среднее 5×5 и Гаусс 5×5 с σ=1. Сверху А, снизу Б.')
p('Для измерения ошибки результаты усреднения сопоставлены с чистыми эталонами. Используются показатели (14): меньшая MSE и большая PSNR означают более точное совпадение яркостей. Контурные карты так с исходным изображением не сравниваются, поскольку представляют другой вид сигнала.')
eq('MSE = (1/(WH)) Σ_x Σ_y (g(x,y)−f(x,y))²;\nPSNR = 10 log₁₀(255²/MSE), дБ.  (14)')
rows=[]
for key,title in [('noisy','Шумный вход'),('mean3','Среднее 3×3'),('mean5','Среднее 5×5'),('gauss5','Гаусс 5×5')]:
    a=metrics['shapes']['methods'][key];b=metrics['detail']['methods'][key]
    rows.append([title,f"{a['mse']:.2f}",f"{a['psnr']:.2f}",f"{b['mse']:.2f}",f"{b['psnr']:.2f}"])
table(['Метод','А: MSE','А: PSNR, дБ','Б: MSE','Б: PSNR, дБ'],rows)
p('На крупных фигурах все усредняющие методы уменьшают ошибку относительно шумного входа. Наименьшую MSE среди проверенных настроек даёт Гаусс 5×5: 95.38 вместо 321.50 у входа. Небольшие веса удалённых соседей ограничивают смешивание разных областей по сравнению с равномерным окном 5×5.')
p('Увеличение равномерного окна с 3×3 до 5×5 делает фон более гладким, но MSE на сцене А возрастает с 111.66 до 147.53. Причиной становится размытие границ. На сцене Б размеры тонких линий сопоставимы с окном, и все три настройки повышают общую MSE относительно шумного входа. Снижение зернистости не компенсирует потерю полезных элементов.')
p('Следовательно, результат необходимо оценивать одновременно по ровным участкам и деталям. Из трёх исследованных настроек гауссово окно 5×5 при σ=1 имеет меньшую ошибку на обеих сценах, однако это наблюдение не доказывает его превосходство при произвольных изображениях и параметрах шума.')

h('Заключение')
p('В работе выполнены два выбранных варианта: выделение контуров и пространственное подавление гауссова шума усреднением. Реализован заданный класс Image, а функции обработки отделены от хранения изображения. Четыре демонстрационных запуска сформировали 28 выходных PNG, проверка из 18 тестовых случаев завершилась успешно.')
p('Операторы Лапласа, Робертса и Собела выявляют перепады яркости, но различаются формой отклика и реакцией на шум. Усредняющие фильтры ослабляют колебания фона ценой размытия. Большое окно не гарантирует меньшую ошибку: сохранность мелких элементов зависит от масштаба окрестности. Иллюстрации и численные оценки подтверждают эти свойства и согласуются с приведёнными формулами.')

# Append only the methods relevant to each group, as in the supplied report.
filters_source=(ROOT/'src/Filters.cpp').read_text()
def function_source(name):
    import re
    match=re.search(r'(?m)^(?:static )?(?:Image|void|int) '+re.escape(name)+r'\([^)]*\)\s*\{',filters_source)
    if not match:raise ValueError('Missing function: '+name)
    opening=filters_source.index('{',match.start())
    depth=0
    for position in range(opening,len(filters_source)):
        if filters_source[position]=='{':depth+=1
        elif filters_source[position]=='}':
            depth-=1
            if depth==0:return filters_source[match.start():position+1]
    raise ValueError('Unbalanced function: '+name)
def listing(text):
    for line in text.splitlines():
        paragraph=doc.add_paragraph()
        run=paragraph.add_run(line)
        run.font.name='Courier New'
h('Приложение А')
p('Фрагменты Filters.cpp для выделения контуров. Сначала приведены общие вспомогательные функции и преобразование яркости, затем расчёт трёх операторов. Общие функции также используются листингом приложения Б. Полный проект содержит необходимые объявления и точку входа.')
common=filters_source[:filters_source.index('namespace Filters {')]
listing(common+'namespace Filters {\n\n'+'\n\n'.join(function_source(name) for name in [
    'toGrayscale','edges','applySobel','applyRoberts','applyLaplacian'])+'\n}\n')
h('Приложение Б')
p('Фрагменты Filters.cpp для усредняющих фильтров. Внутренний расчёт average используется функциями applyBoxBlur и applyGaussianBlur; toGrayscale, sample и put приведены в приложении А.')
listing('namespace Filters {\n\n'+'\n\n'.join(function_source(name) for name in [
    'average','applyBoxBlur','applyGaussianBlur'])+'\n}\n')

def apply_sample_format(document):
    """Typography measured from the supplied classmate report; no text is copied."""
    from docx.shared import Twips
    for section in document.sections:
        section.page_width=Twips(11906)
        section.page_height=Twips(16838)
        section.top_margin=Twips(1134)
        section.bottom_margin=Twips(1134)
        section.left_margin=Twips(1701)
        section.right_margin=Twips(850)
        section.header_distance=Twips(720)
        section.footer_distance=Twips(720)
        section.gutter=Twips(0)
        section.different_first_page_header_footer=True

    def font(run, name='Times New Roman', size=14, bold=False):
        run.font.name=name
        run.font.size=Pt(size)
        run.font.color.rgb=RGBColor(0,0,0)
        run.font.bold=bold
        run.font.italic=False
        run.font.underline=False
        fonts=run._r.get_or_add_rPr().get_or_add_rFonts()
        for key in list(fonts.attrib):
            if 'theme' in key.lower(): del fonts.attrib[key]
        for key in ['ascii','hAnsi','eastAsia','cs']: fonts.set(qn('w:'+key),name)

    normal=document.styles['Normal']
    normal.font.name='Times New Roman'
    normal.font.size=Pt(14)
    normal.font.bold=False
    normal.paragraph_format.first_line_indent=Twips(720)
    normal.paragraph_format.line_spacing=1.5
    normal.paragraph_format.space_before=Pt(0)
    normal.paragraph_format.space_after=Pt(0)
    normal.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    heading=document.styles['Heading 1']
    heading.font.name='Times New Roman'
    heading.font.size=Pt(14)
    heading.font.bold=True
    heading.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
    heading.paragraph_format.first_line_indent=Twips(0)
    heading.paragraph_format.space_before=Pt(0)
    heading.paragraph_format.space_after=Pt(0)
    heading.paragraph_format.line_spacing=1.5

    for i, paragraph in enumerate(document.paragraphs):
        pf=paragraph.paragraph_format
        is_code=any(run.font.name=='Courier New' for run in paragraph.runs)
        is_heading=paragraph.style.name.startswith('Heading')
        is_subheading=paragraph.text in {
            'Оператор Собела','Оператор Робертса','Оператор Лапласа',
            'Прямоугольный усредняющий фильтр (Box Blur)','Фильтр Гаусса (Gaussian Blur)'}
        is_algorithm=paragraph.text.startswith('Алгоритм ')
        is_picture=bool(paragraph._p.xpath('.//w:drawing'))
        is_caption=paragraph.text.startswith('Рисунок ')
        is_equation=paragraph.alignment==WD_ALIGN_PARAGRAPH.CENTER and not is_heading and not is_caption and i>10
        pf.space_before=Pt(0)
        pf.space_after=Pt(0)
        pf.line_spacing=1.5
        pf.first_line_indent=Twips(720)
        pf.left_indent=Twips(0)
        pf.right_indent=Twips(0)
        pf.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
        pf.keep_together=False
        pf.keep_with_next=False
        # Old page breaks were positioned for 10–12 pt text. Let the larger
        # sample typography flow naturally, retaining title/appendix boundaries.
        for page_break in list(paragraph._p.xpath('.//w:br[@w:type="page"]')):
            page_break.getparent().remove(page_break)
        pf.page_break_before=(i==11 or paragraph.text in {
            'Математическое описание алгоритма','Программное описание алгоритма',
            'Результат работы программы','Заключение','Приложение А','Приложение Б'})
        for run in paragraph.runs:font(run,bold=is_heading or is_subheading)
        if is_heading or is_subheading:
            pf.alignment=WD_ALIGN_PARAGRAPH.CENTER
            pf.first_line_indent=Twips(0)
            pf.keep_with_next=True
            pf.keep_together=True
        elif is_code:
            pf.line_spacing=1.0
            pf.first_line_indent=Twips(0)
            pf.widow_control=False
            for run in paragraph.runs:font(run,'Consolas',12)
        elif is_picture or is_caption or is_equation:
            pf.alignment=WD_ALIGN_PARAGRAPH.CENTER
            pf.first_line_indent=Twips(720) if is_caption else Twips(0)
            pf.keep_together=True
            pf.keep_with_next=is_picture
        if is_algorithm:
            pf.first_line_indent=Twips(0)
            pf.keep_with_next=True
            for run in paragraph.runs:run.font.italic=True
        if not paragraph.text and not is_picture:
            pf.line_spacing=Pt(1)
            pf.line_spacing_rule=WD_LINE_SPACING.EXACTLY
            for run in paragraph.runs:font(run,size=1)
        if i<10:
            pf.alignment=WD_ALIGN_PARAGRAPH.CENTER
            pf.first_line_indent=Twips(0)
            pf.space_before=Pt(14) if paragraph.text.strip() and i else Pt(0)
            pf.space_after=Pt(14) if paragraph.text.strip() else Pt(0)
            if not paragraph.text.strip():
                pf.line_spacing=Pt(14)
                pf.line_spacing_rule=WD_LINE_SPACING.EXACTLY
            if i==0:pf.first_line_indent=Twips(720)
            if i==3:
                for run in paragraph.runs:font(run,bold=True)
            if i==7:pf.alignment=WD_ALIGN_PARAGRAPH.LEFT

    for shape in document.inline_shapes:
        if shape.width > Twips(9355):
            shape.height=round(shape.height*Twips(9355)/shape.width)
            shape.width=Twips(9355)
    for table in document.tables:
        table.autofit=False
        column_width=Twips(9355//len(table.columns))
        for column in table.columns:column.width=column_width
        width=table._tbl.tblPr.find(qn('w:tblW'))
        width.set(qn('w:type'),'dxa')
        width.set(qn('w:w'),'9355')
        for row_index,row in enumerate(table.rows):
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            if row_index==0:row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
            for cell in row.cells:
                cell.width=column_width
                for paragraph in cell.paragraphs:
                    pf=paragraph.paragraph_format
                    pf.space_before=pf.space_after=Pt(0)
                    pf.line_spacing=1.5
                    pf.first_line_indent=Twips(0)
                    pf.alignment=WD_ALIGN_PARAGRAPH.LEFT
                    pf.keep_with_next=(row_index==0)
                    for run in paragraph.runs:font(run)
    for section in document.sections:
        for paragraph in section.footer.paragraphs:
            paragraph.paragraph_format.line_spacing=1.15
            paragraph.paragraph_format.first_line_indent=Twips(0)
            paragraph.paragraph_format.space_before=paragraph.paragraph_format.space_after=Pt(0)
            paragraph.alignment=WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:font(run)

apply_sample_format(doc)
# Keep wrapped code statements together and attach closing braces to their body.
for i,paragraph in enumerate(doc.paragraphs):
    if any(run.font.name=='Consolas' for run in paragraph.runs):
        paragraph.paragraph_format.keep_together=True
        if i+1<len(doc.paragraphs):
            following=doc.paragraphs[i+1]
            if any(run.font.name=='Consolas' for run in following.runs):
                paragraph.paragraph_format.keep_with_next=following.text.lstrip().startswith('}')
doc.core_properties.title='Лабораторная работа №2. Варианты 1 и 2'
doc.core_properties.author=''
doc.save(OUT/'Отчет.docx')
office=shutil.which('soffice') or shutil.which('libreoffice')
if not office:raise RuntimeError('LibreOffice is required for PDF export')
with tempfile.TemporaryDirectory(prefix='kg-report-') as directory:
    path=Path(directory)
    subprocess.run([office,'-env:UserInstallation='+(path/'profile').as_uri(),'--headless',
                    '--convert-to','pdf:writer_pdf_Export','--outdir',str(path),str(OUT/'Отчет.docx')],check=True)
    exported=path/'Отчет.pdf'
    if not exported.is_file():raise RuntimeError('PDF export failed')
    shutil.copyfile(exported,OUT/'Отчет.pdf')
print('Created the common report for variants 1 and 2.')
