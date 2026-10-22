"""Export the defense notes as a plain black DOCX and PDF."""
from pathlib import Path
import re
import shutil
import subprocess
import tempfile
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
folder=ROOT/'preparation'
base=folder/'Подготовка_к_защите'
text=base.with_suffix('.md').read_text(encoding='utf-8')
doc=Document()
s=doc.sections[0]
s.page_width=Cm(21)
s.page_height=Cm(29.7)
s.top_margin=s.bottom_margin=Cm(2)
s.left_margin=Cm(3)
s.right_margin=Cm(1.5)
s.footer_distance=Cm(1.25)
for style in doc.styles:
    if hasattr(style,'font'):
        style.font.name='Times New Roman'
        style.font.color.rgb=RGBColor(0,0,0)
        style.font.underline=False
        rpr=style.element.find(qn('w:rPr'))
        if rpr is not None:
            fonts=rpr.find(qn('w:rFonts'))
            if fonts is not None:
                for key in list(fonts.attrib):
                    if 'theme' in key.lower():del fonts.attrib[key]
                for key in ['ascii','hAnsi','eastAsia','cs']:
                    fonts.set(qn('w:'+key),'Times New Roman')
for color in doc.styles.element.iter(qn('w:color')):
    color.attrib.clear()
    color.set(qn('w:val'),'000000')
normal=doc.styles['Normal']
normal.font.size=Pt(14)
normal.paragraph_format.line_spacing=1.5
normal.paragraph_format.first_line_indent=Cm(1.27)
normal.paragraph_format.space_after=Pt(0)
normal.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
for name in ['Heading 1','Heading 2','Heading 3']:
    style=doc.styles[name]
    style.font.size=Pt(14)
    style.font.bold=True
    style.paragraph_format.first_line_indent=Cm(0)
    style.paragraph_format.line_spacing=1.15
    style.paragraph_format.space_before=Pt(10)
    style.paragraph_format.space_after=Pt(4)
    style.paragraph_format.keep_with_next=True
    style.paragraph_format.keep_together=True
footer=s.footer.paragraphs[0]
footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
footer.paragraph_format.first_line_indent=Cm(0)
field=OxmlElement('w:fldSimple')
field.set(qn('w:instr'),'PAGE')
footer._p.append(field)

lines=text.splitlines()
i=0
while i<len(lines):
    line=lines[i].strip()
    if not line:
        i+=1
        continue
    if line.startswith('|'):
        rows=[]
        while i<len(lines) and lines[i].strip().startswith('|'):
            parts=[cell.strip() for cell in lines[i].strip().strip('|').split('|')]
            if not all(re.fullmatch(r'[:\- ]+',cell) for cell in parts):rows.append(parts)
            i+=1
        table=doc.add_table(rows=0,cols=len(rows[0]))
        table.style='Table Grid'
        table.autofit=False
        for column in table.columns:column.width=Cm(16.5/len(rows[0]))
        for row_index,values in enumerate(rows):
            row=table.add_row()
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            if row_index==0:row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
            for cell,value in zip(row.cells,values):
                cell.text=value
                p=cell.paragraphs[0]
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.first_line_indent=Cm(0)
                p.paragraph_format.line_spacing=1.15
                p.paragraph_format.keep_with_next=(row_index==0)
                for run in p.runs:
                    run.font.size=Pt(12)
                    run.font.bold=(row_index==0)
        continue
    if line.startswith('# '):
        p=doc.add_heading(line[2:],level=1)
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    elif line.startswith('## '):
        doc.add_heading(line[3:],level=2)
    elif line.startswith('### '):
        doc.add_heading(line[4:],level=3)
    else:
        p=doc.add_paragraph(line)
        if re.match(r'^\d+\. ',line):
            p.paragraph_format.first_line_indent=Cm(0)
            p.paragraph_format.space_after=Pt(5)
            p.alignment=WD_ALIGN_PARAGRAPH.LEFT
    i+=1

doc.core_properties.title='Подготовка к защите — простыми словами'
doc.core_properties.subject='Разбор программы лабораторной работы, варианты 1–3'
doc.core_properties.author=''
doc.save(base.with_suffix('.docx'))
office=shutil.which('soffice') or shutil.which('libreoffice')
if not office:raise RuntimeError('LibreOffice is required for PDF export')
with tempfile.TemporaryDirectory(prefix='kg-preparation-') as temporary:
    path=Path(temporary)
    subprocess.run([office,'-env:UserInstallation='+(path/'profile').as_uri(),
                    '--headless','--convert-to','pdf:writer_pdf_Export',
                    '--outdir',str(path),str(base.with_suffix('.docx'))],check=True)
    exported=path/(base.name+'.pdf')
    if not exported.is_file():raise RuntimeError('PDF export failed')
    shutil.copyfile(exported,base.with_suffix('.pdf'))
print('Created preparation DOCX and PDF')
